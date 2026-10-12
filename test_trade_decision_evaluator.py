import json
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

from trade_decision_evaluator import evaluate_holdings, evaluate_workbook, write_outputs


def holding(day, isin, quantity, price, *, close=None):
    return {
        "date": day.isoformat(),
        "isin": isin,
        "name": isin,
        "quantity": float(quantity),
        "market_value_base": float(quantity * price),
        "evaluation_pnl_base": 0.0,
        "close_local": float(price if close is None else close),
        "nav": 1_000_000.0,
        "action_weight_source": None,
        "ticker": isin,
        "strategy_code": "TEST",
    }


class TradeDecisionEvaluatorTest(unittest.TestCase):
    def test_weekday_quantity_or_composition_change_is_session_but_weekend_is_not(self):
        thursday = date(2024, 1, 4)
        friday = date(2024, 1, 5)
        saturday = date(2024, 1, 6)
        monday = date(2024, 1, 8)
        tuesday = date(2024, 1, 9)
        holdings = [
            holding(thursday, "A", 10, 100),
            holding(friday, "A", 15, 100),
            holding(saturday, "A", 20, 100),
            holding(monday, "A", 20, 101),
            holding(tuesday, "A", 20, 101),
            holding(tuesday, "B", 5, 50),
        ]

        result = evaluate_holdings(holdings, bootstrap_samples=10)
        actions = {(row["date"], row["isin"]): row for row in result["actions"]}

        self.assertEqual(actions[(friday.isoformat(), "A")]["evaluation_quality_flag"], "OK")
        self.assertEqual(
            actions[(friday.isoformat(), "A")]["future_date_1"], monday.isoformat()
        )
        self.assertEqual(
            actions[(saturday.isoformat(), "A")]["evaluation_quality_flag"],
            "NON_SESSION_DATE",
        )
        self.assertEqual(actions[(tuesday.isoformat(), "B")]["evaluation_quality_flag"], "OK")

    def test_session_proxy_price_priority_and_previous_close_exit(self):
        thursday = date(2024, 1, 4)
        friday = date(2024, 1, 5)
        saturday = date(2024, 1, 6)
        monday = date(2024, 1, 8)
        tuesday = date(2024, 1, 9)
        holdings = [
            holding(thursday, "A", 10, 99),
            holding(thursday, "B", 10, 200),
            holding(thursday, "C", 10, 100),
            holding(friday, "A", 15, 100),
            holding(friday, "B", 10, 201),
            holding(friday, "C", 20, 50),
            holding(saturday, "A", 20, 100),
            holding(saturday, "B", 10, 201),
            holding(saturday, "C", 20, 50),
            holding(monday, "A", 25, 101),
            holding(monday, "C", 20, 51),
            holding(tuesday, "A", 25, 102),
            holding(tuesday, "B", 0, 190, close=190),
            holding(tuesday, "C", 20, 52),
        ]

        result = evaluate_holdings(holdings, cost_bps=10, bootstrap_samples=20)
        actions = {(row["date"], row["isin"]): row for row in result["actions"]}
        friday_buy = actions[(friday.isoformat(), "A")]
        weekend_buy = actions[(saturday.isoformat(), "A")]
        exit_action = actions[(monday.isoformat(), "B")]
        corporate_action = actions[(friday.isoformat(), "C")]

        self.assertEqual(friday_buy["future_date_1"], monday.isoformat())
        self.assertEqual(friday_buy["base_price_source"], "VALUATION_UNIT_PROXY")
        self.assertEqual(weekend_buy["evaluation_quality_flag"], "NON_SESSION_DATE")
        self.assertFalse(weekend_buy["statistics_included"])
        self.assertEqual(corporate_action["source_quality_flag"], "POSSIBLE_CORPORATE_ACTION")
        self.assertFalse(corporate_action["statistics_included"])
        self.assertEqual(exit_action["base_date_source"], "PREVIOUS_CLOSE_PROXY")
        self.assertEqual(exit_action["future_price_source_1"], "LOCAL_CLOSE_FALLBACK")
        self.assertAlmostEqual(exit_action["gross_directional_return_1"], -(190 / 201 - 1))

        total_h1 = next(
            row for row in result["summary"] if row["scope"] == "TOTAL" and row["horizon"] == 1
        )
        self.assertEqual(total_h1["excluded_quality_actions"], 2)
        total_h20 = next(
            row for row in result["summary"] if row["scope"] == "TOTAL" and row["horizon"] == 20
        )
        self.assertEqual(total_h20["complete"], 0)
        self.assertEqual(total_h20["coverage"], 0)

    def test_positive_primary_verdict_with_enough_days(self):
        days = []
        current = date(2024, 1, 1)
        while len(days) < 65:
            if current.weekday() < 5:
                days.append(current)
            current += timedelta(days=1)
        holdings = [
            holding(day, "UP", index + 1, 100 * (1.01**index))
            for index, day in enumerate(days)
        ]

        result = evaluate_holdings(holdings, cost_bps=10, bootstrap_samples=100)
        primary = next(
            row for row in result["summary"] if row["scope"] == "TOTAL" and row["horizon"] == 5
        )
        self.assertGreaterEqual(primary["coverage"], 0.8)
        self.assertGreaterEqual(primary["unique_days"], 30)
        self.assertGreaterEqual(primary["ess_days"], 20)
        self.assertEqual(primary["primary_verdict"], "POSITIVE")

    def test_workbook_pipeline_and_four_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "ss.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            headers = [
                "기준일", "자산구분", "종목코드", "종목명", "액면/주수",
                "평가금액(기준)", "평가손익(기준)", "당일종가", "순자산", "NAV비율", "전략코드",
            ]
            for column, header in enumerate(headers, 1):
                sheet.cell(7, column, header)
            rows = [
                [datetime(2024, 1, 4), "주식", "ISIN1", "Sample", 10, 990, 0, 99, 10000, 9.9, "S1"],
                [datetime(2024, 1, 5), "주식", "ISIN1", "Sample", 15, 1500, 0, 100, 10000, 15, "S1"],
                [datetime(2024, 1, 8), "주식", "ISIN1", "Sample", 15, 1515, 0, 101, 10000, 15, "S1"],
            ]
            for row_number, values in enumerate(rows, 8):
                for column, value in enumerate(values, 1):
                    sheet.cell(row_number, column, value)
            mapping = workbook.create_sheet("ISIN Map")
            mapping.append(["ISIN", "Ticker"])
            mapping.append(["ISIN1", "TEST US Equity"])
            workbook.save(workbook_path)

            result = evaluate_workbook(workbook_path, bootstrap_samples=10)
            paths = write_outputs(result, root / "report.json")

            self.assertEqual(set(result), {"metadata", "summary", "daily", "actions"})
            self.assertEqual(set(paths), {"json", "summary_csv", "daily_csv", "actions_csv"})
            self.assertTrue(all(path.exists() for path in paths.values()))
            loaded = json.loads(paths["json"].read_text(encoding="utf-8"))
            self.assertIn("not realized P&L", loaded["metadata"]["sell_markout_interpretation"])


if __name__ == "__main__":
    unittest.main()
