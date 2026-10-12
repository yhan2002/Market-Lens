import sqlite3
import tempfile
import unittest
import csv
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

import openpyxl

from trading_log import (
    import_trading_log,
    backfill_holdings_csv,
    infer_actions,
    latest_long_tickers,
    read_holdings_workbook,
    read_holdings_csv,
)


class TradingLogTest(unittest.TestCase):
    def test_import_replaces_snapshots_with_optional_history_preservation(self) -> None:
        def holding(day: str, isin: str, quantity: float, ticker: str) -> dict[str, object]:
            return {
                "date": day, "isin": isin, "ticker": ticker, "name": ticker,
                "quantity": quantity, "market_value_base": quantity * 10,
                "evaluation_pnl_base": 0, "close_local": 10, "nav": 1000,
                "action_weight_source": None, "strategy_code": None,
            }

        original = [
            holding("2025-03-01", "ISIN1", 10, "AAA"),
            holding("2025-03-01", "ISIN2", 5, "EXIT"),
            holding("2025-03-03", "ISIN1", 20, "AAA"),
            holding("2025-03-03", "ISIN2", 5, "EXIT"),
            holding("2025-03-05", "ISIN1", 30, "AAA"),
        ]
        original_map = {
            "ISIN1": ("AAA US Equity", "AAA"),
            "ISIN2": ("EXIT US Equity", "EXIT"),
            "LEGACY": ("OLD US Equity", "OLD"),
        }
        incoming = [holding("2025-03-03", "ISIN1", 15, "NEW")]
        incoming_map = {"ISIN1": ("NEW US Equity", "NEW")}

        for preserve in (False, True):
            with self.subTest(preserve_history=preserve), tempfile.TemporaryDirectory() as directory:
                database = Path(directory) / "router.db"
                workbook = Path(directory) / "mocked.xlsx"
                with closing(sqlite3.connect(database)) as connection:
                    connection.executescript(
                        """
                        CREATE TABLE metadata(key TEXT PRIMARY KEY, value TEXT NOT NULL);
                        CREATE TABLE universe(
                            ticker TEXT, group_name TEXT, source TEXT, active INTEGER,
                            PRIMARY KEY(ticker, group_name)
                        );
                        """
                    )
                with patch("trading_log.read_holdings_workbook", side_effect=[
                    (original, original_map), (incoming, incoming_map),
                ]):
                    import_trading_log(database, workbook)
                    if preserve:
                        counts = import_trading_log(database, workbook, preserve_history=True)
                    else:
                        counts = import_trading_log(database, workbook)

                expected = [original[0], original[1], incoming[0], original[4]] if preserve else incoming
                with closing(sqlite3.connect(database)) as connection:
                    connection.row_factory = sqlite3.Row
                    rows = [dict(row) for row in connection.execute(
                        "SELECT * FROM daily_holdings ORDER BY date, isin"
                    )]
                    actions = [tuple(row) for row in connection.execute(
                        """
                        SELECT date, previous_date, isin, action_type, previous_quantity, new_quantity
                        FROM inferred_actions ORDER BY date, isin
                        """
                    )]
                    mappings = {
                        row["isin"]: (row["bloomberg_ticker"], row["yahoo_ticker"])
                        for row in connection.execute("SELECT * FROM ticker_map")
                    }
                    active = [row[0] for row in connection.execute(
                        "SELECT ticker FROM universe WHERE group_name='TRADING_LOG' AND active=1"
                    )]
                self.assertEqual(rows, expected)
                self.assertEqual(mappings, original_map | incoming_map if preserve else incoming_map)
                self.assertEqual(active, ["AAA"] if preserve else ["NEW"])
                self.assertEqual(counts, (4, 3) if preserve else (1, 0))
                self.assertEqual(actions, [
                    ("2025-03-03", "2025-03-01", "ISIN1", "INCREASE", 10, 15),
                    ("2025-03-03", "2025-03-01", "ISIN2", "EXIT", 5, 0),
                    ("2025-03-05", "2025-03-03", "ISIN1", "INCREASE", 15, 30),
                ] if preserve else [])

    def test_csv_backfill_adds_only_missing_equity_rows(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "holdings.csv"
            with source.open("w", encoding="utf-8-sig", newline="") as output:
                writer = csv.writer(output)
                writer.writerows([
                    ["header", "ISIN", "quantity", "ticker"],
                    ["2025-03-12", "CASH", "100", "USD Curncy"],
                    ["2025-03-12", "ISIN1", "2", "AAA US Equity"],
                    ["2025-03-12", "ISIN1", "3", "AAA US Equity"],
                    ["2025-03-13", "ISIN1", "6", "AAA US Equity"],
                ])
            holdings, _ = read_holdings_csv(source)
            self.assertEqual([row["quantity"] for row in holdings], [5.0, 6.0])

            database = root / "router.db"
            with closing(sqlite3.connect(database)) as connection:
                connection.executescript(
                    """
                    CREATE TABLE ticker_map(isin TEXT PRIMARY KEY, bloomberg_ticker TEXT, yahoo_ticker TEXT);
                    CREATE TABLE daily_holdings(
                        date TEXT, isin TEXT, ticker TEXT, name TEXT, quantity REAL,
                        market_value_base REAL, evaluation_pnl_base REAL, close_local REAL,
                        nav REAL, action_weight_source REAL, strategy_code TEXT,
                        PRIMARY KEY(date, isin)
                    );
                    INSERT INTO daily_holdings VALUES
                        ('2025-03-12','ISIN1','AAA','AAA',5,NULL,NULL,NULL,NULL,NULL,NULL);
                    """
                )
            dry_run = backfill_holdings_csv(database, source)
            self.assertEqual((dry_run["missing_rows"], dry_run["conflict_rows"]), (1, 0))
            applied = backfill_holdings_csv(database, source, apply=True)
            self.assertTrue(applied["applied"])
            with closing(sqlite3.connect(database)) as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM daily_holdings").fetchone()[0], 2)

    def test_net_quantity_changes_become_actions(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "log.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            headers = [
                "기준일", "자산구분", "종목코드", "종목명", "액면/주수",
                "평가금액(기준)", "평가손익(기준)", "당일종가", "순자산", "NAV비율", "전략코드",
            ]
            for column, header in enumerate(headers, 1):
                sheet.cell(7, column, header)
            rows = [
                ["2024-01-02", "주식", "ISIN1", "Sample", 10, 1000, 0, 100, 10000, 10, "S1"],
                ["2024-01-03", "주식", "ISIN1", "Sample", 15, 1650, 50, 110, 10100, 16.3, "S1"],
                ["2024-01-04", "주식", "ISIN1", "Sample", 0, 0, 0, 105, 10200, 0, "S1"],
                ["2024-01-04", "주식", "ISIN2", "Second", 5, 500, 0, 100, 10200, 5, "S1"],
            ]
            from datetime import datetime
            for row_number, values in enumerate(rows, 8):
                values[0] = datetime.fromisoformat(values[0])
                for column, value in enumerate(values, 1):
                    sheet.cell(row_number, column, value)
            mapping = workbook.create_sheet("ISIN Map")
            mapping.append(["ISIN", "Ticker"])
            mapping.append(["ISIN1", "TEST US Equity"])
            mapping.append(["ISIN2", "SECOND US Equity"])
            workbook.save(path)

            holdings, _ = read_holdings_workbook(path)
            actions = infer_actions(holdings)
            self.assertEqual(
                [action["action_type"] for action in actions],
                ["INCREASE", "EXIT", "ENTRY"],
            )
            self.assertEqual([action["side"] for action in actions], ["BUY", "SELL", "BUY"])
            self.assertAlmostEqual(actions[0]["action_weight"], 0.0544554455)

            database = Path(directory) / "router.db"
            with closing(sqlite3.connect(database)) as connection:
                connection.executescript(
                    """
                    CREATE TABLE metadata(key TEXT PRIMARY KEY, value TEXT NOT NULL);
                    CREATE TABLE universe(
                        ticker TEXT, group_name TEXT, source TEXT, active INTEGER,
                        PRIMARY KEY(ticker, group_name)
                    );
                    """
                )
                connection.commit()
            import_trading_log(database, path)
            self.assertEqual(latest_long_tickers(database), ("SECOND",))
            with closing(sqlite3.connect(database)) as connection:
                active = connection.execute(
                    "SELECT ticker FROM universe WHERE group_name='TRADING_LOG' AND active=1"
                ).fetchall()
            self.assertEqual(active, [("SECOND",)])


if __name__ == "__main__":
    unittest.main()
