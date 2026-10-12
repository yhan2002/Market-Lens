from __future__ import annotations

import argparse
import csv
import sqlite3
from collections import defaultdict
from contextlib import closing
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl


TRADING_SCHEMA = """
CREATE TABLE IF NOT EXISTS ticker_map (
    isin TEXT PRIMARY KEY,
    bloomberg_ticker TEXT,
    yahoo_ticker TEXT
);
CREATE TABLE IF NOT EXISTS daily_holdings (
    date TEXT NOT NULL,
    isin TEXT NOT NULL,
    ticker TEXT,
    name TEXT,
    quantity REAL NOT NULL,
    market_value_base REAL,
    evaluation_pnl_base REAL,
    close_local REAL,
    nav REAL,
    action_weight_source REAL,
    strategy_code TEXT,
    PRIMARY KEY (date, isin)
);
CREATE TABLE IF NOT EXISTS inferred_actions (
    action_id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    previous_date TEXT NOT NULL,
    isin TEXT NOT NULL,
    ticker TEXT,
    action_type TEXT NOT NULL,
    side TEXT NOT NULL,
    previous_quantity REAL NOT NULL,
    new_quantity REAL NOT NULL,
    delta_quantity REAL NOT NULL,
    action_notional_base REAL,
    nav REAL,
    action_weight REAL,
    strategy_code TEXT,
    quality_flag TEXT NOT NULL,
    UNIQUE(date, isin)
);
CREATE INDEX IF NOT EXISTS idx_inferred_actions_date ON inferred_actions(date);
"""


def _repair_text(value: object) -> object:
    if not isinstance(value, str):
        return value
    try:
        return value.encode("gbk").decode("cp949")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return value


def _yahoo_ticker(bloomberg_ticker: object) -> str | None:
    if not isinstance(bloomberg_ticker, str) or not bloomberg_ticker.strip():
        return None
    parts = bloomberg_ticker.strip().upper().split()
    symbol = parts[0]
    market = parts[1] if len(parts) > 1 else ""
    if market == "HK" and symbol.isdigit():
        return f"{symbol.zfill(4)}.HK"
    if market == "JP" and symbol.isdigit():
        return f"{symbol}.T"
    return symbol if market in {"", "US"} else None


def _number(value: object) -> float | None:
    return float(value) if isinstance(value, (int, float)) else None


def _strategy(values: set[str]) -> str | None:
    cleaned = sorted(value for value in values if value)
    return "|".join(cleaned) if cleaned else None


def read_holdings_workbook(workbook_path: Path) -> tuple[list[dict[str, object]], dict[str, tuple[str, str | None]]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        mapping_sheet = workbook.worksheets[-1]
        ticker_map = {
            str(row[0]).strip(): (str(row[1]).strip(), _yahoo_ticker(row[1]))
            for row in mapping_sheet.iter_rows(min_row=2, values_only=True)
            if row[0] and row[1]
        }

        sheet = workbook["Sheet1"]
        headers = [_repair_text(cell.value) for cell in next(sheet.iter_rows(min_row=7, max_row=7))]
        columns = {header: index for index, header in enumerate(headers)}
        required = {
            "기준일", "자산구분", "종목코드", "종목명", "액면/주수",
            "평가금액(기준)", "평가손익(기준)", "당일종가", "순자산", "NAV비율", "전략코드",
        }
        missing = required - set(columns)
        if missing:
            raise ValueError(f"Sheet1 missing columns: {sorted(missing)}")

        aggregated: dict[tuple[str, str], dict[str, object]] = {}
        for row in sheet.iter_rows(min_row=8, values_only=True):
            current_date = row[columns["기준일"]]
            asset_type = _repair_text(row[columns["자산구분"]])
            isin = row[columns["종목코드"]]
            quantity = _number(row[columns["액면/주수"]])
            if not isinstance(current_date, (datetime, date)) or asset_type != "주식" or not isin or quantity is None:
                continue

            day = current_date.date().isoformat() if isinstance(current_date, datetime) else current_date.isoformat()
            key = (day, str(isin).strip())
            record = aggregated.setdefault(
                key,
                {
                    "date": day,
                    "isin": key[1],
                    "name": str(_repair_text(row[columns["종목명"]]) or ""),
                    "quantity": 0.0,
                    "market_value_base": 0.0,
                    "evaluation_pnl_base": 0.0,
                    "close_local": _number(row[columns["당일종가"]]),
                    "nav": _number(row[columns["순자산"]]),
                    "action_weight_source": _number(row[columns["NAV비율"]]),
                    "strategy_codes": set(),
                },
            )
            record["quantity"] = float(record["quantity"]) + quantity
            for target, source in (("market_value_base", "평가금액(기준)"), ("evaluation_pnl_base", "평가손익(기준)")):
                value = _number(row[columns[source]])
                if value is not None:
                    record[target] = float(record[target]) + value
            strategy_code = row[columns["전략코드"]]
            if strategy_code not in (None, ""):
                record["strategy_codes"].add(str(strategy_code))

        holdings = []
        for record in aggregated.values():
            bloomberg, yahoo = ticker_map.get(str(record["isin"]), (None, None))
            record["ticker"] = yahoo
            record["strategy_code"] = _strategy(record.pop("strategy_codes"))
            holdings.append(record)
        return sorted(holdings, key=lambda row: (str(row["date"]), str(row["isin"]))), ticker_map
    finally:
        workbook.close()


def read_holdings_csv(csv_path: Path) -> tuple[list[dict[str, object]], dict[str, tuple[str, str | None]]]:
    aggregated: dict[tuple[str, str], dict[str, object]] = {}
    ticker_map: dict[str, tuple[str, str | None]] = {}
    with csv_path.open("r", encoding="utf-8-sig", newline="") as source:
        for row_number, row in enumerate(csv.reader(source), 1):
            if len(row) != 4:
                continue
            try:
                day = date.fromisoformat(row[0].strip()).isoformat()
            except ValueError:
                continue
            bloomberg = row[3].strip()
            if not bloomberg.upper().endswith(" EQUITY"):
                continue
            isin = row[1].strip()
            yahoo = _yahoo_ticker(bloomberg)
            if not isin or yahoo is None:
                raise ValueError(f"CSV row {row_number} has an unmapped equity")
            try:
                quantity = float(row[2].replace(",", ""))
            except ValueError as error:
                raise ValueError(f"CSV row {row_number} has an invalid quantity") from error
            if quantity <= 0:
                raise ValueError(f"CSV row {row_number} has a non-positive quantity")
            previous = ticker_map.setdefault(isin, (bloomberg, yahoo))
            if previous != (bloomberg, yahoo):
                raise ValueError(f"CSV maps {isin} to multiple tickers")
            key = (day, isin)
            record = aggregated.setdefault(
                key,
                {
                    "date": day,
                    "isin": isin,
                    "ticker": yahoo,
                    "name": bloomberg.split()[0],
                    "quantity": 0.0,
                    "market_value_base": None,
                    "evaluation_pnl_base": None,
                    "close_local": None,
                    "nav": None,
                    "action_weight_source": None,
                    "strategy_code": None,
                },
            )
            record["quantity"] = float(record["quantity"]) + quantity
    return sorted(aggregated.values(), key=lambda row: (str(row["date"]), str(row["isin"]))), ticker_map


def backfill_holdings_csv(database: Path, csv_path: Path, *, apply: bool = False) -> dict[str, object]:
    holdings, ticker_map = read_holdings_csv(csv_path)
    if not holdings:
        raise ValueError("CSV has no mapped equity holdings")
    incoming = {(str(row["date"]), str(row["isin"])): row for row in holdings}
    connection_target = database if apply else f"file:{database.resolve().as_posix()}?mode=ro"
    with closing(sqlite3.connect(connection_target, uri=not apply)) as connection:
        columns = {
            str(row[1]) for row in connection.execute("PRAGMA table_info(daily_holdings)")
        }
        if not {"date", "isin", "ticker", "quantity"}.issubset(columns):
            raise ValueError("database daily_holdings schema is incompatible")
        existing = {
            (str(day)[:10], str(isin)): float(quantity)
            for day, isin, quantity in connection.execute(
                "SELECT date, isin, quantity FROM daily_holdings"
            )
        }
        overlap = set(incoming) & set(existing)
        conflicts = sorted(
            key
            for key in overlap
            if abs(existing[key] - float(incoming[key]["quantity"])) > 1e-9
        )
        missing = [incoming[key] for key in sorted(set(incoming) - set(existing))]
        if apply:
            if conflicts:
                raise ValueError(f"CSV conflicts with {len(conflicts)} existing holdings rows")
            connection.executemany(
                """
                INSERT INTO ticker_map(isin, bloomberg_ticker, yahoo_ticker) VALUES (?, ?, ?)
                ON CONFLICT(isin) DO UPDATE SET
                    bloomberg_ticker=excluded.bloomberg_ticker,
                    yahoo_ticker=excluded.yahoo_ticker
                """,
                [(isin, bloomberg, yahoo) for isin, (bloomberg, yahoo) in ticker_map.items()],
            )
            connection.executemany(
                "INSERT INTO daily_holdings VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                [
                    tuple(row[key] for key in (
                        "date", "isin", "ticker", "name", "quantity",
                        "market_value_base", "evaluation_pnl_base", "close_local",
                        "nav", "action_weight_source", "strategy_code",
                    ))
                    for row in missing
                ],
            )
            connection.commit()
    missing_dates = sorted({str(row["date"]) for row in missing})
    return {
        "source_rows": len(holdings),
        "source_dates": len({str(row["date"]) for row in holdings}),
        "source_start": min(str(row["date"]) for row in holdings),
        "source_end": max(str(row["date"]) for row in holdings),
        "overlap_rows": len(overlap),
        "conflict_rows": len(conflicts),
        "missing_rows": len(missing),
        "missing_dates": len(missing_dates),
        "missing_start": missing_dates[0] if missing_dates else None,
        "missing_end": missing_dates[-1] if missing_dates else None,
        "applied": apply,
    }


def _unit_value(record: dict[str, object] | None) -> float | None:
    if not record:
        return None
    quantity = float(record["quantity"])
    value = _number(record["market_value_base"])
    return abs(value / quantity) if value is not None and quantity else None


def _quality_flag(previous: dict[str, object] | None, current: dict[str, object] | None, ticker: str | None) -> str:
    if ticker is None:
        return "MISSING_TICKER"
    if not previous or not current:
        return "OK"
    previous_quantity = abs(float(previous["quantity"]))
    current_quantity = abs(float(current["quantity"]))
    previous_unit = _unit_value(previous)
    current_unit = _unit_value(current)
    if not previous_quantity or not current_quantity or not previous_unit or not current_unit:
        return "OK"
    quantity_ratio = current_quantity / previous_quantity
    common_ratios = (0.1, 0.2, 0.25, 0.5, 2.0, 3.0, 4.0, 5.0, 10.0)
    if min(abs(quantity_ratio - ratio) / ratio for ratio in common_ratios) < 0.01:
        if abs(quantity_ratio * current_unit / previous_unit - 1) < 0.08:
            return "POSSIBLE_CORPORATE_ACTION"
    return "OK"


def infer_actions(holdings: list[dict[str, object]]) -> list[dict[str, object]]:
    snapshots: dict[str, dict[str, dict[str, object]]] = defaultdict(dict)
    for holding in holdings:
        snapshots[str(holding["date"])][str(holding["isin"])] = holding
    dates = sorted(snapshots)
    actions: list[dict[str, object]] = []

    for previous_date, current_date in zip(dates, dates[1:]):
        previous_snapshot = snapshots[previous_date]
        current_snapshot = snapshots[current_date]
        for isin in sorted(set(previous_snapshot) | set(current_snapshot)):
            previous = previous_snapshot.get(isin)
            current = current_snapshot.get(isin)
            previous_quantity = float(previous["quantity"]) if previous else 0.0
            current_quantity = float(current["quantity"]) if current else 0.0
            delta = current_quantity - previous_quantity
            if abs(delta) < 1e-9:
                continue

            if previous_quantity == 0:
                action_type = "ENTRY"
            elif current_quantity == 0:
                action_type = "EXIT"
            elif abs(current_quantity) > abs(previous_quantity):
                action_type = "INCREASE"
            else:
                action_type = "DECREASE"
            source = current if delta > 0 else previous
            assert source is not None
            unit_value = _unit_value(source)
            notional = abs(delta) * unit_value if unit_value is not None else None
            nav = _number((current or previous).get("nav"))
            ticker = str(source["ticker"]) if source.get("ticker") else None
            actions.append(
                {
                    "date": current_date,
                    "previous_date": previous_date,
                    "isin": isin,
                    "ticker": ticker,
                    "action_type": action_type,
                    "side": "BUY" if delta > 0 else "SELL",
                    "previous_quantity": previous_quantity,
                    "new_quantity": current_quantity,
                    "delta_quantity": delta,
                    "action_notional_base": notional,
                    "nav": nav,
                    "action_weight": abs(notional / nav) if notional is not None and nav else None,
                    "strategy_code": source.get("strategy_code"),
                    "quality_flag": _quality_flag(previous, current, ticker),
                }
            )
    return actions


def import_trading_log(
    database: Path, workbook_path: Path, *, preserve_history: bool = False,
) -> tuple[int, int]:
    """Replace workbook snapshots; optionally retain other dates and return merged counts."""
    holdings, ticker_map = read_holdings_workbook(workbook_path)
    if not holdings:
        raise ValueError("workbook has no equity holdings")
    with closing(sqlite3.connect(database)) as connection:
        connection.executescript(TRADING_SCHEMA)
        connection.execute("DELETE FROM inferred_actions")
        if preserve_history:
            connection.executemany(
                "DELETE FROM daily_holdings WHERE date=?",
                [(day,) for day in sorted({str(row["date"]) for row in holdings})],
            )
        else:
            connection.execute("DELETE FROM daily_holdings")
            connection.execute("DELETE FROM ticker_map")
        connection.executemany(
            """
            INSERT INTO ticker_map(isin, bloomberg_ticker, yahoo_ticker) VALUES (?, ?, ?)
            ON CONFLICT(isin) DO UPDATE SET
                bloomberg_ticker=excluded.bloomberg_ticker,
                yahoo_ticker=excluded.yahoo_ticker
            """,
            [(isin, bloomberg, yahoo) for isin, (bloomberg, yahoo) in ticker_map.items()],
        )
        connection.executemany(
            """
            INSERT INTO daily_holdings VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    row["date"], row["isin"], row["ticker"], row["name"], row["quantity"],
                    row["market_value_base"], row["evaluation_pnl_base"], row["close_local"],
                    row["nav"], row["action_weight_source"], row["strategy_code"],
                )
                for row in holdings
            ],
        )
        if preserve_history:
            connection.row_factory = sqlite3.Row
            holdings = [
                dict(row) for row in connection.execute(
                    "SELECT * FROM daily_holdings ORDER BY date, isin"
                )
            ]
        actions = infer_actions(holdings)
        connection.executemany(
            """
            INSERT INTO inferred_actions(
                date, previous_date, isin, ticker, action_type, side,
                previous_quantity, new_quantity, delta_quantity,
                action_notional_base, nav, action_weight, strategy_code, quality_flag
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [tuple(row[key] for key in (
                "date", "previous_date", "isin", "ticker", "action_type", "side",
                "previous_quantity", "new_quantity", "delta_quantity", "action_notional_base",
                "nav", "action_weight", "strategy_code", "quality_flag",
            )) for row in actions],
        )
        latest_date = max(str(row["date"]) for row in holdings)
        yahoo_tickers = sorted(
            {
                row["ticker"]
                for row in holdings
                if row["ticker"]
                and str(row["date"]) == latest_date
                and float(row["quantity"]) > 0
            }
        )
        connection.execute("UPDATE universe SET active=0 WHERE group_name='TRADING_LOG'")
        connection.executemany(
            """
            INSERT INTO universe(ticker, group_name, source, active)
            VALUES (?, 'TRADING_LOG', ?, 1)
            ON CONFLICT(ticker, group_name) DO UPDATE SET source=excluded.source, active=1
            """,
            [(ticker, str(workbook_path)) for ticker in yahoo_tickers],
        )
        connection.execute(
            "INSERT OR REPLACE INTO metadata(key, value) VALUES ('trading_log_imported_at', ?)",
            (datetime.now(timezone.utc).isoformat(),),
        )
        connection.commit()
    return len(holdings), len(actions)


def latest_long_tickers(database: Path, as_of: str | None = None) -> tuple[str, ...]:
    where = "WHERE date<=?" if as_of else ""
    parameters = (as_of,) if as_of else ()
    with closing(sqlite3.connect(database)) as connection:
        latest = connection.execute(
            f"SELECT MAX(date) FROM daily_holdings {where}", parameters
        ).fetchone()
        if not latest or not latest[0]:
            return ()
        rows = connection.execute(
            """
            SELECT ticker
            FROM daily_holdings
            WHERE date=? AND ticker IS NOT NULL AND quantity>0
            GROUP BY ticker
            ORDER BY ticker
            """,
            (latest[0],),
        ).fetchall()
    return tuple(str(row[0]) for row in rows)


def print_summary(database: Path) -> None:
    with closing(sqlite3.connect(database)) as connection:
        rows = connection.execute(
            """
            SELECT action_type, side, quality_flag, COUNT(1)
            FROM inferred_actions
            GROUP BY action_type, side, quality_flag
            ORDER BY action_type, side, quality_flag
            """
        ).fetchall()
        coverage = connection.execute(
            """
            SELECT COUNT(DISTINCT a.ticker),
                   COUNT(DISTINCT CASE WHEN p.ticker IS NOT NULL THEN a.ticker END)
            FROM inferred_actions a
            LEFT JOIN daily_prices p ON p.ticker=a.ticker
            """
        ).fetchone()
    for row in rows:
        print(" | ".join(map(str, row)))
    print(f"price_coverage: {coverage[1]}/{coverage[0]} tickers")


def main() -> None:
    parser = argparse.ArgumentParser(description="Net-holdings decision log")
    commands = parser.add_subparsers(dest="command", required=True)
    importer = commands.add_parser("import-xlsx", help="import holdings and infer net actions")
    importer.add_argument("database", type=Path)
    importer.add_argument("workbook", type=Path)
    csv_backfill = commands.add_parser("backfill-csv", help="append missing equity holdings from CSV")
    csv_backfill.add_argument("database", type=Path)
    csv_backfill.add_argument("csv", type=Path)
    csv_backfill.add_argument("--apply", action="store_true")
    summary = commands.add_parser("summary", help="show inferred action counts")
    summary.add_argument("database", type=Path)
    args = parser.parse_args()

    if args.command == "import-xlsx":
        holdings, actions = import_trading_log(args.database, args.workbook)
        print(f"imported {holdings} daily holdings and inferred {actions} net actions")
    elif args.command == "backfill-csv":
        summary = backfill_holdings_csv(args.database, args.csv, apply=args.apply)
        print(" | ".join(f"{key}={value}" for key, value in summary.items()))
    else:
        print_summary(args.database)


if __name__ == "__main__":
    main()
